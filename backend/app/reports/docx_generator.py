"""
Phase 21: DOCX Report Generator
Uses python-docx for editable Word documents
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from typing import Dict, Any, List
from datetime import datetime


class DOCXReportGenerator:
    """Generates professional DOCX reports"""
    
    def __init__(self):
        self.doc = Document()
        self._setup_styles()
    
    def _safe_float(self, value, default: float = 0.0) -> float:
        """Safely convert any value to float"""
        if value is None:
            return default
        try:
            return float(value)
        except (ValueError, TypeError):
            return default
    
    def _setup_styles(self):
        """Setup document styles"""
        style = self.doc.styles['Normal']
        style.font.name = 'Calibri'
        style.font.size = Pt(11)
    
    def generate(self, analysis: Dict[str, Any], output_path: str):
        """Generate DOCX report"""
        self.doc = Document()
        self._setup_styles()
        
        self._build_cover(analysis)
        self.doc.add_page_break()
        
        self._build_executive_summary(analysis)
        self.doc.add_page_break()
        
        self._build_features(analysis)
        self._build_team(analysis)
        self._build_tasks(analysis)
        self._build_cost(analysis)
        self._build_timeline(analysis)
        self._build_risks(analysis)
        self._build_mvp(analysis)
        self._build_tech(analysis)
        self._build_assumptions(analysis)
        
        self.doc.save(output_path)
    
    def _add_heading(self, text: str, level: int = 1, color: str = None):
        """Add heading with optional color"""
        heading = self.doc.add_heading(text, level=level)
        if color:
            for run in heading.runs:
                run.font.color.rgb = RGBColor.from_string(color)
        return heading
    
    # ============================================
    # COVER PAGE
    # ============================================
    
    def _build_cover(self, analysis: Dict):
        """Cover page"""
        project = analysis.get("project") or {}
        
        # Title
        title = self.doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("PROJECT SCOPE REPORT")
        run.font.size = Pt(28)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1e, 0x3a, 0x8a)
        
        self.doc.add_paragraph()
        
        # Project name
        name_p = self.doc.add_paragraph()
        name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        name_run = name_p.add_run(project.get("name", "Project"))
        name_run.font.size = Pt(22)
        name_run.font.bold = True
        
        self.doc.add_paragraph()
        self.doc.add_paragraph()
        
        # Meta table
        meta_table = self.doc.add_table(rows=6, cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # ✅ Safe float conversions
        total_hours = self._safe_float(analysis.get('total_estimated_hours', 0))
        total_cost = self._safe_float(
            (analysis.get('cost') or {}).get('total', {}).get('expected', 0)
        )
        
        meta = [
            ("Generated", datetime.now().strftime("%B %d, %Y")),
            ("Platform", str(project.get("platform", "Web"))),
            ("Status", str(project.get("status", "Draft"))),
            ("Total Hours", f"{total_hours:.0f}"),
            ("Total Cost", f"${total_cost:,.2f}"),
            ("Risk Level", str(analysis.get('risk_level', 'MEDIUM'))),
        ]
        
        for i, (k, v) in enumerate(meta):
            meta_table.rows[i].cells[0].text = k
            meta_table.rows[i].cells[1].text = str(v)
        
        self.doc.add_paragraph()
        self.doc.add_paragraph()
        
        # Description
        desc = project.get("description", "")
        if desc:
            self._add_heading("Project Description", 2)
            self.doc.add_paragraph(str(desc)[:800])
    
    # ============================================
    # EXECUTIVE SUMMARY
    # ============================================
    
    def _build_executive_summary(self, analysis: Dict):
        """Executive summary"""
        self._add_heading("Executive Summary", 1, "1e40af")
        
        explanation = analysis.get("explanation") or {}
        summary = explanation.get("summary", "")
        if summary:
            self.doc.add_paragraph(str(summary))
        
        self.doc.add_paragraph()
        self._add_heading("Key Metrics", 2)
        
        cost = (analysis.get("cost") or {}).get("total", {}) or {}
        timeline = analysis.get("timeline") or {}
        
        cost_expected = self._safe_float(cost.get('expected', 0))
        cost_min = self._safe_float(cost.get('min', 0))
        cost_max = self._safe_float(cost.get('max', 0))
        total_hours = self._safe_float(analysis.get('total_estimated_hours', 0))
        complexity = self._safe_float(analysis.get('complexity_score', 0))
        
        metrics = [
            ("Metric", "Value"),
            ("Total Estimated Hours", f"{total_hours:.0f} hours"),
            ("Total Cost (Expected)", f"${cost_expected:,.2f}"),
            ("Total Cost Range", f"${cost_min:,.2f} - ${cost_max:,.2f}"),
            ("Timeline", f"{timeline.get('total_working_days', 0)} working days"),
            ("Complexity Score", f"{complexity:.1f}/100"),
            ("Risk Level", str(analysis.get('risk_level', 'N/A'))),
            ("Features", str(len(analysis.get('features', [])))),
            ("Tasks", str(len(analysis.get('tasks', [])))),
        ]
        
        table = self.doc.add_table(rows=len(metrics), cols=2)
        table.style = 'Light Grid Accent 1'
        for i, (k, v) in enumerate(metrics):
            table.rows[i].cells[0].text = k
            table.rows[i].cells[1].text = str(v)
    
    # ============================================
    # FEATURES
    # ============================================
    
    def _build_features(self, analysis: Dict):
        """Features section"""
        self._add_heading("Requirements & Features", 1, "1e40af")
        
        features = analysis.get("features", []) or []
        if not features:
            return
        
        self._add_heading(f"Extracted Features ({len(features)})", 2)
        
        table = self.doc.add_table(rows=1, cols=3)
        table.style = 'Light Grid Accent 1'
        headers = table.rows[0].cells
        headers[0].text = "Feature"
        headers[1].text = "Priority"
        headers[2].text = "Complexity"
        
        for f in features[:25]:
            row = table.add_row().cells
            row[0].text = str(f.get("canonical_name", ""))
            row[1].text = str(f.get("priority", ""))
            row[2].text = f"{self._safe_float(f.get('complexity', 0)):.0f}/10"
        
        self.doc.add_paragraph()
    
    # ============================================
    # TEAM
    # ============================================
    
    def _build_team(self, analysis: Dict):
        """Team allocation"""
        self._add_heading("Team Allocation", 1, "1e40af")
        
        summary = analysis.get("summary", {}) or {}
        if not summary:
            return
        
        table = self.doc.add_table(rows=1, cols=4)
        table.style = 'Light Grid Accent 1'
        headers = table.rows[0].cells
        headers[0].text = "Role"
        headers[1].text = "Hours"
        headers[2].text = "Tasks"
        headers[3].text = "Cost"
        
        for role, data in summary.items():
            row = table.add_row().cells
            row[0].text = str(role)
            row[1].text = f"{self._safe_float(data.get('total_hours', 0)):.0f}"
            row[2].text = str(data.get('num_tasks', 0))
            row[3].text = f"${self._safe_float(data.get('estimated_cost', 0)):,.2f}"
        
        self.doc.add_paragraph()
    
    # ============================================
    # TASKS
    # ============================================
    
    def _build_tasks(self, analysis: Dict):
        """Task breakdown"""
        self._add_heading("Task Breakdown", 1, "1e40af")
        
        tasks = analysis.get("tasks", []) or []
        if not tasks:
            return
        
        self.doc.add_paragraph(f"Total Tasks: {len(tasks)}")
        
        table = self.doc.add_table(rows=1, cols=4)
        table.style = 'Light Grid Accent 1'
        headers = table.rows[0].cells
        headers[0].text = "Task"
        headers[1].text = "Role"
        headers[2].text = "Hours"
        headers[3].text = "Priority"
        
        for task in tasks[:30]:
            row = table.add_row().cells
            row[0].text = str(task.get("title") or "")[:60]
            row[1].text = f"Role {task.get('role_id', 0)}"
            row[2].text = f"{self._safe_float(task.get('estimated_hours', 0)):.0f}"
            row[3].text = str(task.get("priority", ""))
        
        self.doc.add_paragraph()
    
    # ============================================
    # COST
    # ============================================
    
    def _build_cost(self, analysis: Dict):
        """Cost section"""
        self.doc.add_page_break()
        self._add_heading("Cost Estimation", 1, "1e40af")
        
        cost = analysis.get("cost") or {}
        total = cost.get("total", {}) or {}
        
        table = self.doc.add_table(rows=4, cols=2)
        table.style = 'Light Grid Accent 1'
        
        rows = [
            ("Scenario", "Amount"),
            ("Best Case (Min)", f"${self._safe_float(total.get('min', 0)):,.2f}"),
            ("Expected", f"${self._safe_float(total.get('expected', 0)):,.2f}"),
            ("Worst Case (Max)", f"${self._safe_float(total.get('max', 0)):,.2f}"),
        ]
        
        for i, (k, v) in enumerate(rows):
            table.rows[i].cells[0].text = k
            table.rows[i].cells[1].text = v
        
        self.doc.add_paragraph()
        
        # By role
        by_role = cost.get("by_role", {}) or {}
        if by_role:
            self._add_heading("Cost by Role", 2)
            table2 = self.doc.add_table(rows=1, cols=4)
            table2.style = 'Light Grid Accent 1'
            h = table2.rows[0].cells
            h[0].text = "Role"
            h[1].text = "Hours"
            h[2].text = "Rate"
            h[3].text = "Cost"
            
            for role, data in by_role.items():
                cv = data.get("cost", 0)
                if isinstance(cv, dict):
                    cv = cv.get("expected", 0)
                cv_float = self._safe_float(cv)
                hours = self._safe_float(data.get('hours', 0))
                rate = self._safe_float(data.get('rate', 0))
                
                row = table2.add_row().cells
                row[0].text = str(role)
                row[1].text = f"{hours:.0f}"
                row[2].text = f"${rate:.0f}/hr"
                row[3].text = f"${cv_float:,.2f}"
    
    # ============================================
    # TIMELINE
    # ============================================
    
    def _build_timeline(self, analysis: Dict):
        """Timeline section"""
        self.doc.add_page_break()
        self._add_heading("Timeline & Critical Path", 1, "1e40af")
        
        timeline = analysis.get("timeline") or {}
        if not timeline:
            return
        
        self.doc.add_paragraph(
            f"Start: {(timeline.get('start_date') or '')[:10]}  |  "
            f"End: {(timeline.get('end_date') or '')[:10]}  |  "
            f"Duration: {timeline.get('total_working_days', 0)} working days"
        )
        
        milestones = timeline.get("milestones", []) or []
        if milestones:
            self._add_heading("Milestones", 2)
            table = self.doc.add_table(rows=1, cols=3)
            table.style = 'Light Grid Accent 1'
            h = table.rows[0].cells
            h[0].text = "Milestone"
            h[1].text = "Date"
            h[2].text = "Description"
            
            for m in milestones:
                row = table.add_row().cells
                row[0].text = str(m.get("name", ""))
                row[1].text = (m.get("date") or "")[:10]
                row[2].text = str(m.get("description") or "")[:60]
    
    # ============================================
    # RISKS
    # ============================================
    
    def _build_risks(self, analysis: Dict):
        """Risk section"""
        self._add_heading("Risk Assessment", 1, "1e40af")
        
        risks = analysis.get("risks") or {}
        if not risks:
            return
        
        risk_score = self._safe_float(risks.get('risk_score', 0))
        self.doc.add_paragraph(
            f"Total Risks: {risks.get('total_risks', 0)} | "
            f"Risk Level: {risks.get('risk_level', 'N/A')} | "
            f"Risk Score: {risk_score:.1f}"
        )
        
        top = risks.get("top_risks", []) or []
        if top:
            table = self.doc.add_table(rows=1, cols=4)
            table.style = 'Light Grid Accent 1'
            h = table.rows[0].cells
            h[0].text = "Risk"
            h[1].text = "Level"
            h[2].text = "Impact"
            h[3].text = "Mitigation"
            
            for r in top:
                row = table.add_row().cells
                row[0].text = str(r.get("name") or "")[:35]
                row[1].text = str(r.get("risk_level", ""))
                row[2].text = str(r.get("impact", ""))
                row[3].text = str(r.get("mitigation") or "")[:50]
        
        recs = risks.get("recommendations", []) or []
        if recs:
            self._add_heading("Recommendations", 2)
            for r in recs:
                self.doc.add_paragraph(f"• {r}")
    
    # ============================================
    # MVP
    # ============================================
    
    def _build_mvp(self, analysis: Dict):
        """MVP section"""
        self._add_heading("MVP Recommendation", 1, "1e40af")
        
        features = analysis.get("features", []) or []
        high = [f for f in features if str(f.get("priority") or "").upper() in ["HIGH", "CRITICAL"]]
        
        self.doc.add_paragraph(
            f"MVP should focus on {len(high)} high-priority features. "
            f"Remaining {len(features) - len(high)} features can be Phase 2."
        )
        
        for f in high[:10]:
            self.doc.add_paragraph(
                f"• {f.get('canonical_name', '')} — {str(f.get('description') or '')[:80]}",
                style='List Bullet'
            )
    
    # ============================================
    # TECHNOLOGY
    # ============================================
    
    def _build_tech(self, analysis: Dict):
        """Technology section"""
        self._add_heading("Technology Recommendations", 1, "1e40af")
        
        features = [f.get("canonical_name", "") for f in analysis.get("features", []) or []]
        
        stack = [
            ("Frontend", "Next.js 14 + TypeScript + Tailwind CSS"),
            ("Backend", "Python FastAPI + Pydantic v2"),
            ("Database", "PostgreSQL 16"),
            ("Cache/Queue", "Redis + Celery"),
            ("Deployment", "Docker + GitHub Actions"),
        ]
        
        if "AUTHENTICATION" in features:
            stack.append(("Auth", "JWT + bcrypt/argon2"))
        if "PAYMENT" in features:
            stack.append(("Payments", "Stripe / PayPal"))
        if "REAL_TIME" in features:
            stack.append(("Real-time", "WebSockets + Redis Pub/Sub"))
        if "MOBILE_APP" in features:
            stack.append(("Mobile", "React Native / Flutter"))
        
        table = self.doc.add_table(rows=len(stack), cols=2)
        table.style = 'Light Grid Accent 1'
        for i, (k, v) in enumerate(stack):
            table.rows[i].cells[0].text = k
            table.rows[i].cells[1].text = v
    
    # ============================================
    # ASSUMPTIONS
    # ============================================
    
    def _build_assumptions(self, analysis: Dict):
        """Assumptions and limitations"""
        self.doc.add_page_break()
        self._add_heading("Assumptions & Limitations", 1, "1e40af")
        
        explanation = analysis.get("explanation") or {}
        
        assumptions = explanation.get("assumptions", []) or []
        if assumptions:
            self._add_heading("Assumptions", 2)
            for a in assumptions:
                self.doc.add_paragraph(f"• {a}", style='List Bullet')
        
        limitations = explanation.get("limitations", []) or []
        if limitations:
            self._add_heading("Limitations", 2)
            for l in limitations:
                self.doc.add_paragraph(f"• {l}", style='List Bullet')
        
        hybrid = analysis.get("hybrid_estimate") or {}
        confidence = self._safe_float(hybrid.get("confidence", 0))
        self._add_heading("Confidence Level", 2)
        self.doc.add_paragraph(
            f"Overall confidence: {confidence * 100:.0f}%. "
            f"Based on data quality, model performance, and project complexity."
        )